import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


source = Path(sys.argv[1])
target = Path(sys.argv[2])
target.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(source, target)

document = Document(target)

for table_number, table in enumerate(document.tables, 1):
    if table_number in (1, 2):
        continue

    target_size = 10 if table_number == 20 else 10.5
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if not run.text:
                        continue
                    run.font.size = Pt(target_size)
                    if table_number == 20:
                        run.font.name = "Consolas"
                        rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                        rfonts.set(qn("w:ascii"), "Consolas")
                        rfonts.set(qn("w:hAnsi"), "Consolas")
                        rfonts.set(qn("w:eastAsia"), "Consolas")

document.save(target)
print(target)

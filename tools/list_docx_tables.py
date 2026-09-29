import json
import sys
from collections import Counter
from docx import Document


doc = Document(sys.argv[1])
result = []
for index, table in enumerate(doc.tables, 1):
    sizes = Counter()
    texts = []
    for row in table.rows:
        for cell in row.cells:
            value = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            if value:
                texts.append(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        size = run.font.size.pt if run.font.size else None
                        sizes["inherited" if size is None else f"{size:g}"] += len(run.text)
    result.append({
        "table": index,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "sizes_by_chars": dict(sizes),
        "sample": " | ".join(texts)[:220],
    })
print(json.dumps(result, ensure_ascii=False, indent=2))

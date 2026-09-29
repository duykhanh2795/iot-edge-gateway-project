import json
import sys
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile
from lxml import etree

from docx import Document
from docx.oxml.ns import qn


def pt(value):
    return None if value is None else round(value.pt, 2)


def font_name(run):
    if run.font.name:
        return run.font.name
    rpr = run._element.rPr
    if rpr is not None and rpr.rFonts is not None:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            value = rpr.rFonts.get(qn(f"w:{key}"))
            if value:
                return value
    return None


def style_info(style):
    pf = style.paragraph_format
    font = style.font
    return {
        "name": style.name,
        "base": style.base_style.name if style.base_style else None,
        "font": font.name,
        "size_pt": pt(font.size),
        "bold": font.bold,
        "italic": font.italic,
        "alignment": str(pf.alignment),
        "space_before_pt": pt(pf.space_before),
        "space_after_pt": pt(pf.space_after),
        "line_spacing": str(pf.line_spacing),
        "first_line_indent_cm": None if pf.first_line_indent is None else round(pf.first_line_indent.cm, 3),
    }


def inspect(path):
    doc = Document(path)
    style_counts = Counter()
    run_fonts = Counter()
    run_sizes = Counter()
    by_style = defaultdict(lambda: {"paragraphs": 0, "fonts": Counter(), "sizes": Counter()})

    def scan_paragraph(paragraph):
        style_name = paragraph.style.name if paragraph.style else "(none)"
        if paragraph.text.strip():
            style_counts[style_name] += 1
            by_style[style_name]["paragraphs"] += 1
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            name = font_name(run) or "(inherited)"
            size = pt(run.font.size)
            size_key = "(inherited)" if size is None else str(size)
            run_fonts[name] += len(run.text)
            run_sizes[size_key] += len(run.text)
            by_style[style_name]["fonts"][name] += len(run.text)
            by_style[style_name]["sizes"][size_key] += len(run.text)

    body_sizes = Counter()
    table_sizes = Counter()

    def collect_sizes(paragraph, counter):
        for run in paragraph.runs:
            if run.text.strip():
                size = pt(run.font.size)
                counter["(inherited)" if size is None else str(size)] += len(run.text)

    for paragraph in doc.paragraphs:
        scan_paragraph(paragraph)
        collect_sizes(paragraph, body_sizes)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    scan_paragraph(paragraph)
                    collect_sizes(paragraph, table_sizes)

    package_defaults = {}
    with ZipFile(path) as package:
        styles_xml = etree.fromstring(package.read("word/styles.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        rpr = styles_xml.find(".//w:docDefaults/w:rPrDefault/w:rPr", ns)
        ppr = styles_xml.find(".//w:docDefaults/w:pPrDefault/w:pPr", ns)
        if rpr is not None:
            fonts = rpr.find("w:rFonts", ns)
            size = rpr.find("w:sz", ns)
            package_defaults["run_fonts"] = dict(fonts.attrib) if fonts is not None else {}
            package_defaults["run_size_half_points"] = size.get(qn("w:val")) if size is not None else None
        if ppr is not None:
            spacing = ppr.find("w:spacing", ns)
            package_defaults["paragraph_spacing"] = dict(spacing.attrib) if spacing is not None else {}

    sections = []
    for section in doc.sections:
        sections.append({
            "width_cm": round(section.page_width.cm, 3),
            "height_cm": round(section.page_height.cm, 3),
            "top_cm": round(section.top_margin.cm, 3),
            "bottom_cm": round(section.bottom_margin.cm, 3),
            "left_cm": round(section.left_margin.cm, 3),
            "right_cm": round(section.right_margin.cm, 3),
            "header_cm": round(section.header_distance.cm, 3),
            "footer_cm": round(section.footer_distance.cm, 3),
        })

    wanted = ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]
    styles = {}
    for name in wanted:
        try:
            styles[name] = style_info(doc.styles[name])
        except KeyError:
            pass

    return {
        "path": str(Path(path).resolve()),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": sections,
        "style_counts": dict(style_counts),
        "direct_run_fonts_by_chars": dict(run_fonts),
        "direct_run_sizes_by_chars": dict(run_sizes),
        "body_run_sizes_by_chars": dict(body_sizes),
        "table_run_sizes_by_chars": dict(table_sizes),
        "package_defaults": package_defaults,
        "styles": styles,
        "usage_by_style": {
            name: {
                "paragraphs": data["paragraphs"],
                "fonts_by_chars": dict(data["fonts"]),
                "sizes_by_chars": dict(data["sizes"]),
            }
            for name, data in by_style.items()
        },
    }


if __name__ == "__main__":
    print(json.dumps(inspect(sys.argv[1]), ensure_ascii=False, indent=2))

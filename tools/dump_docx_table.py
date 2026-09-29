import json
import sys

from docx import Document


document = Document(sys.argv[1])
for requested in (int(value) for value in sys.argv[2:]):
    table = document.tables[requested - 1]
    print(f"TABLE {requested}")
    rows = []
    for row in table.rows:
        rows.append(["\n".join(p.text for p in cell.paragraphs) for cell in row.cells])
    print(json.dumps(rows, ensure_ascii=False, indent=2))

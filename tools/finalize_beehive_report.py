from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"F:\iot-edge-gateway-project\output\midterm-beehive\Bao_cao_IoT_HTX_Ong_Mat_Cao_Nguyen_Xanh_font_chuan_mau.docx")
OUTPUT = Path(r"F:\iot-edge-gateway-project\output\midterm-beehive\Bao_cao_IoT_HTX_Ong_Mat_Cao_Nguyen_Xanh_hoan_thien.docx")
DIAGRAM = Path(r"F:\iot-edge-gateway-project\output\midterm-beehive\kien_truc_he_thong_to_ong_dong_bo.png")


def set_font(run, size=10.5, bold=None, name="Times New Roman"):
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_text(cell, text, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, code=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    set_font(run, size=size, name="Consolas" if code else "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    for old in tr_pr.findall(qn("w:tblHeader")):
        tr_pr.remove(old)
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def replace_body(table, rows, alignments, size=10.5):
    template = deepcopy(table.rows[1]._tr)
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    repeat_header(table.rows[0])
    prevent_split(table.rows[0])
    for values in rows:
        tr = deepcopy(template)
        table._tbl.append(tr)
        row = table.rows[-1]
        prevent_split(row)
        for idx, value in enumerate(values):
            set_cell_text(row.cells[idx], value, size=size, align=alignments[idx])


def enable_field_updates(doc):
    settings = doc.settings.element
    current = settings.find(qn("w:updateFields"))
    if current is None:
        current = OxmlElement("w:updateFields")
        settings.append(current)
    current.set(qn("w:val"), "true")


def replace_diagram(table):
    cell = table.cell(0, 0)
    for paragraph in list(cell.paragraphs):
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(6.05))


EXAMPLE_TOPICS = [
    ("bee/site-a/hive/hive-017/telemetry/snapshot", "telemetry", "Nhiệt độ, độ ẩm, cân nặng, góc nghiêng và pin mỗi 60 giây"),
    ("bee/site-b/hive/hive-031/telemetry/acoustic", "telemetry", "Đặc trưng âm thanh mỗi 120 giây, không chứa audio thô"),
    ("bee/site-c/hive/hive-023/event/tilt", "event", "Thùng ong nghiêng hoặc bị dịch chuyển bất thường"),
    ("bee/site-c/hive/hive-023/event/weight-drop", "event", "Khối lượng giảm nhanh vượt ngưỡng"),
    ("bee/site-a/hive/hive-017/state/status", "state", "Online/offline, pin và thời điểm lastSeen gần nhất"),
    ("bee/site-a/hive/hive-017/state/firmware", "state", "Kết quả OTA và phiên bản firmware mà node đang chạy"),
    ("bee/site-c/weather/weather-01/telemetry/snapshot", "telemetry", "Điều kiện môi trường chung của trại C"),
    ("bee/site-c/gateway/edge-01/state/status", "state", "Trạng thái online/offline và lastSeen của gateway"),
    ("bee/site-c/gateway/edge-01/event/data-degraded", "event", "Gateway báo giảm tần suất gửi do mất 4G hoặc hàng đợi cao"),
    ("bee/site-b/hive/hive-031/cmd/config", "cmd", "Đổi chu kỳ gửi hoặc ngưỡng tại một node"),
    ("bee/site-b/hive/hive-031/cmd/calibrate", "cmd", "Yêu cầu hiệu chuẩn load cell của node"),
    ("bee/site-b/hive/hive-031/cmd/audio-capture", "cmd", "Yêu cầu thu audio 10 giây và lưu tại gateway"),
    ("bee/site-b/hive/hive-031/cmd/restart", "cmd", "Khởi động lại có kiểm soát một node"),
    ("bee/site-b/hive/hive-031/cmd/ota", "cmd", "Yêu cầu node cập nhật đúng phiên bản firmware"),
    ("bee/site-c/alarm/siren-01/cmd/set", "cmd", "Bật hoặc tắt còi tại trại C"),
]


TOPIC_REGISTRY = [
    ("hive/+/telemetry/snapshot", "out", "1", "Không", "Dữ liệu chuỗi thời gian; deviceId + seq giúp khử lặp."),
    ("hive/+/telemetry/acoustic", "out", "1", "Không", "Cần cho phân tích xu hướng, không phải trạng thái tức thời."),
    ("hive/+/event/tilt", "out", "1", "Không", "Không được mất; không retained để tránh phát lại cảnh báo cũ."),
    ("hive/+/event/weight-drop", "out", "1", "Không", "Sự kiện an ninh; eventId giúp xử lý lặp an toàn."),
    ("hive/+/state/status", "out", "1", "Có", "Subscriber mới cần biết ngay trạng thái gần nhất của node."),
    ("hive/+/state/firmware", "out", "1", "Không", "Ghi kết quả OTA/phiên bản; tránh coi kết quả cũ là hiện tại."),
    ("weather/+/telemetry/snapshot", "out", "1", "Không", "Dữ liệu chuỗi thời gian; retained không cần thiết."),
    ("gateway/+/state/status", "out", "1", "Có", "Dịch vụ mới kết nối cần biết ngay trạng thái gateway."),
    ("gateway/+/event/data-degraded", "out", "1", "Không", "Sự kiện giảm chất lượng; không phát lại như cảnh báo mới."),
    ("hive/+/cmd/config", "in", "1", "Không", "Lệnh một lần; commandId/expiresAt ngăn chạy lại."),
    ("hive/+/cmd/calibrate", "in", "1", "Không", "Hiệu chuẩn là thao tác một lần, không tự chạy lại."),
    ("hive/+/cmd/audio-capture", "in", "1", "Không", "Yêu cầu thu âm một lần, có commandId và expiresAt."),
    ("hive/+/cmd/restart", "in", "1", "Không", "Không để node reconnect rồi nhận lại lệnh restart cũ."),
    ("hive/+/cmd/ota", "in", "1", "Không", "Không để node mới kết nối nhận nhầm chiến dịch OTA cũ."),
    ("alarm/+/cmd/set", "in", "1", "Không", "Tránh còi tự bật lại do retained command cũ."),
]


ACL_BRIDGE_READ = "hive/+/telemetry/#; hive/+/event/#; hive/+/state/#; weather/+/telemetry/#; gateway/+/state/#; gateway/+/event/#"
ACL_INGEST_READ = "bee/+/hive/+/telemetry/#; bee/+/hive/+/event/#; bee/+/hive/+/state/#; bee/+/weather/+/telemetry/#; bee/+/gateway/+/state/#; bee/+/gateway/+/event/#"

BRIDGE_CONFIG = '''# Mosquitto cục bộ - ví dụ site-c
persistence true
persistence_location /var/lib/mosquitto/
autosave_interval 60
max_queued_messages 160000
max_queued_bytes 67108864
queue_qos0_messages false

connection central-site-c
address mqtt-center.example:8883
remote_clientid bridge-site-c
bridge_protocol_version mqttv50
cleansession false
start_type automatic
restart_timeout 5 120
try_private true
bridge_cafile /etc/mosquitto/ca.crt
bridge_certfile /etc/mosquitto/site-c.crt
bridge_keyfile /etc/mosquitto/site-c.key
topic hive/+/telemetry/# out 1 "" bee/site-c/
topic hive/+/event/# out 1 "" bee/site-c/
topic hive/+/state/# out 1 "" bee/site-c/
topic weather/+/telemetry/# out 1 "" bee/site-c/
topic gateway/+/state/# out 1 "" bee/site-c/
topic gateway/+/event/# out 1 "" bee/site-c/
topic hive/+/cmd/# in 1 "" bee/site-c/
topic alarm/+/cmd/# in 1 "" bee/site-c/

# ACL node hive-023
user hive-023
topic write hive/hive-023/telemetry/#
topic write hive/hive-023/event/#
topic write hive/hive-023/state/#
topic read  hive/hive-023/cmd/#'''


def main():
    if not SOURCE.exists() or not DIAGRAM.exists():
        raise FileNotFoundError("Thiếu file nguồn hoặc hình kiến trúc")
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    replace_diagram(doc.tables[3])
    replace_body(doc.tables[5], EXAMPLE_TOPICS,
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    replace_body(doc.tables[7], TOPIC_REGISTRY,
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    acl = doc.tables[14]
    set_cell_text(acl.rows[3].cells[2], ACL_BRIDGE_READ)
    set_cell_text(acl.rows[3].cells[3], "Sáu topic filter outbound, chỉ lấy dữ liệu site C.")
    set_cell_text(acl.rows[5].cells[2], ACL_INGEST_READ)
    set_cell_text(acl.rows[5].cells[3], "Sáu topic filter cho ba site, không ghi command.")

    config = doc.tables[19]
    set_cell_text(config.cell(0, 0), BRIDGE_CONFIG, size=10, code=True)

    # Bắt đầu phần kiến trúc ở trang mới để tiêu đề 2.1 không bị mồ côi
    # ở cuối trang trước khi hình lớn được chuyển sang trang kế tiếp.
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "2. Kiến trúc hệ thống":
            paragraph.paragraph_format.page_break_before = True
            break

    enable_field_updates(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
